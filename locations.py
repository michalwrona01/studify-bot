"""Read current workbook addresses and resolve subject/facility pairs."""

import json
import re
import unicodedata
from pathlib import Path


CATALOG = json.loads(Path(__file__).with_name("locations.json").read_text(encoding="utf-8"))["locations"]


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value).casefold().replace("ł", "l"))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("pulmunologii", "pulmonologii").replace("tarakochirurgii", "torakochirurgii")
    return " ".join(re.findall(r"\w+", text))


RULES = [
    ([normalize(subject) for subject in row["subjects"]],
     [normalize(facility) for facility in row["facilities"]], row)
    for row in CATALOG
]


def resolve_location(name, mode="", catalog=None):
    text = normalize(name)
    if re.search(r"\b(online|on line|zdal\w*)\b", text + " " + normalize(mode)):
        return {}
    matches = {}
    rules = RULES if catalog is None else [
        ([normalize(value) for value in row["subjects"]],
         [normalize(value) for value in row["facilities"]], row) for row in catalog
    ]
    for subjects, facilities, row in rules:
        if any(subject in text for subject in subjects) and any(facility in text for facility in facilities):
            matches[(row["address"], row["lat"], row["lng"])] = {
                "address": row["address"], "lat": row["lat"], "lng": row["lng"]
            }
    # A subject alone is insufficient: several subjects use multiple hospitals.
    return next(iter(matches.values())) if len(matches) == 1 else {}


def facility_name(address):
    return re.split(r"(?:,?\s+)(?:ul\.|al\.|aleja\b|plac\b)", address, maxsplit=1, flags=re.I)[0].strip(" ,")


def read_locations(workbook, geocoder):
    """Read ADRESY (or second sheet), including vertically merged subject cells."""
    sheets = workbook.worksheets
    sheet = next((sheet for sheet in sheets if normalize(sheet.title) == "adresy"), None)
    if sheet is None:
        sheet = sheets[1] if len(sheets) > 1 else None
    if sheet is None:
        return []
    columns = None
    subject = None
    rows = []
    for cells in sheet.iter_rows(values_only=True):
        labels = [normalize(value or "") for value in cells]
        if "przedmiot" in labels and "adres placowki" in labels:
            columns = (labels.index("przedmiot"), labels.index("adres placowki"))
            subject = None
            continue
        if columns is None:
            continue
        subject_cell, address_cell = (cells[index] if index < len(cells) else None for index in columns)
        if subject_cell:
            subject = " ".join(str(subject_cell).split())
        if not address_cell or not subject:
            continue
        address = " ".join(str(address_cell).split())
        facility = facility_name(address)
        subjects, facilities = [subject], [facility]
        # Known aliases aid matching, but addresses always come from this workbook.
        for known in CATALOG:
            if normalize(subject) in [normalize(value) for value in known["subjects"]] and normalize(facility) == normalize(facility_name(known["address"])):
                subjects = list(dict.fromkeys(subjects + known["subjects"]))
                facilities = list(dict.fromkeys(facilities + known["facilities"]))
        point = geocoder.lookup(address) or {}
        rows.append({"subjects": subjects, "facilities": facilities, "address": address,
                     "lat": point.get("lat"), "lng": point.get("lng")})
    return rows


def load_workbook_locations(path, cache_path):
    import openpyxl
    from geocoding import Geocoder
    from logger import logger

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    geocoder = Geocoder(cache_path, seed=CATALOG)
    try:
        rows = read_locations(workbook, geocoder)
        if not rows:
            logger.warning("No address rows found in the current workbook")
        return rows
    finally:
        geocoder.session.close()
        workbook.close()
